import csv
import io
import os
import re
import secrets
import smtplib
import sqlite3
import webbrowser
import base64
from datetime import datetime
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
from pathlib import Path

from flask import Flask, flash, jsonify, redirect, render_template, request, send_file, session, url_for
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

BASE_DIR = Path(__file__).resolve().parent
DATA_DIR = BASE_DIR / "data"
REPORT_DIR = BASE_DIR / "reports"
DB_PATH = DATA_DIR / "simulator.db"
DATA_DIR.mkdir(exist_ok=True)
REPORT_DIR.mkdir(exist_ok=True)


def load_dotenv(path: Path) -> None:
    """Small .env loader so the app does not need python-dotenv."""
    if not path.exists():
        return
    for raw in path.read_text(encoding="utf-8").splitlines():
        line = raw.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, value = line.split("=", 1)
        key = key.strip()
        value = value.strip().strip('"').strip("'")
        os.environ.setdefault(key, value)


load_dotenv(BASE_DIR / ".env")

app = Flask(__name__)
app.secret_key = os.environ.get("FLASK_SECRET_KEY") or secrets.token_hex(32)
app.config.update(SESSION_COOKIE_HTTPONLY=True, SESSION_COOKIE_SAMESITE="Lax")

# Public-repository safe design:
# - No team member names, student IDs, or personal email addresses are stored here.
# - SMTP credentials must stay in .env and must never be committed.
# - Submitted passwords are recorded only as [REDACTED].


PROJECT = {
    "name": "Social Engineering Simulator",
    "description": "Implementing Secured Flask Phishing Simulation Campaign For Training Demonstrations",
    "start_date": " ",
    "end_date": " ",
    "status": " ",
    "company": " ",
    "company_email": " ",
}

SCENARIOS = {
    "account-alert": {
        "name": "Account Security Alert",
        "brand": "NexaSecure",
        "icon": "NS",
        "category": "Account security",
        "phishing_type": "Credential Phishing",
        "subject": "Unusual sign-in detected on your account",
        "sender_name": "NexaSecure Account Services",
        "sender": "alerts@nexasecure-support.example",
        "urgency": "ACTION REQUIRED",
        "preview": "We detected a sign-in from a device that is not associated with your recent activity. Review the sign-in details to keep your account active.",
        "cta": "Review Sign-in Activity",
        "headline": "A new sign-in needs your attention",
        "body": [
            "We noticed a recent sign-in attempt from a new device. If this was you, no action is required.",
            "If you do not recognize the activity, review the session immediately and confirm your account details.",
        ],
        "indicators": [
            "Unexpected security alert",
            "Urgent request to review activity",
            "Sender domain is not the normal service domain",
            "Login details requested after following the message link",
        ],
        "login_title": "Confirm account access",
        "login_subtitle": "Sign in to review the recent account activity.",
    },
    "payment": {
        "name": "Payment Verification",
        "brand": "PayWise",
        "icon": "PW",
        "category": "Payments",
        "phishing_type": "Payment Phishing",
        "subject": "Payment verification required",
        "sender_name": "PayWise Billing",
        "sender": "billing@paywise-review.example",
        "urgency": "TIME SENSITIVE",
        "preview": "A recent payment could not be confirmed. Review your billing details to avoid interruption to your current service.",
        "cta": "Review Billing Information",
        "headline": "Your payment needs attention",
        "body": [
            "Our billing system was unable to confirm the payment method associated with your account.",
            "Please review the current billing information before the next processing window.",
        ],
        "indicators": [
            "Unexpected billing notice",
            "Fear of service interruption",
            "Financial information requested through a message link",
            "Time pressure discourages independent verification",
        ],
        "login_title": "Verify your billing account",
        "login_subtitle": "Sign in to continue to the billing portal.",
    },
    "it-support": {
        "name": "IT Support Request",
        "brand": "WorkHub",
        "icon": "IT",
        "category": "Corporate IT",
        "phishing_type": "Spear-Phishing / IT Impersonation",
        "subject": "Remote access verification required",
        "sender_name": "WorkHub IT Service Desk",
        "sender": "helpdesk@workhub-access.example",
        "urgency": "MAINTENANCE WINDOW",
        "preview": "Your remote-access profile is scheduled for an update. Complete the verification before the maintenance window closes.",
        "cta": "Verify Remote Access",
        "headline": "Remote access verification",
        "body": [
            "A routine access maintenance window is scheduled for your account.",
            "Verify your sign-in details before the maintenance window closes to avoid an interruption to remote access.",
        ],
        "indicators": [
            "Message appears to come from internal IT",
            "Maintenance deadline creates pressure",
            "Unexpected sign-in request",
            "External login page is used for a corporate action",
        ],
        "login_title": "Remote access sign-in",
        "login_subtitle": "Verify your work account to continue.",
    },
    "cloud-share": {
        "name": "Cloud Document Share",
        "brand": "CloudSphere",
        "icon": "CS",
        "category": "Cloud services",
        "phishing_type": "Cloud Credential Phishing",
        "subject": "A document has been shared with you",
        "sender_name": "CloudSphere Sharing",
        "sender": "sharing@cloudsphere-docs.example",
        "urgency": "NEW SHARE",
        "preview": "A document named “Quarterly Review” has been shared with your account. Open the document to review the latest version.",
        "cta": "Open Shared Document",
        "headline": "A document is waiting for you",
        "body": [
            "A colleague has shared a document with your account.",
            "Open the secure document viewer below to review the latest version and comments.",
        ],
        "indicators": [
            "Unexpected document share",
            "Curiosity encourages the recipient to open the link",
            "Sender domain differs from the normal service",
            "Login is requested before viewing the document",
        ],
        "login_title": "Open shared document",
        "login_subtitle": "Sign in to access the document viewer.",
    },
    "delivery": {
        "name": "Delivery Notification",
        "brand": "ParcelDesk",
        "icon": "PD",
        "category": "Delivery",
        "phishing_type": "Link-Based Credential Phishing",
        "subject": "Delivery attempt requires confirmation",
        "sender_name": "ParcelDesk Delivery",
        "sender": "dispatch@parceldesk-track.example",
        "urgency": "DELIVERY UPDATE",
        "preview": "We could not complete the delivery attempt. Confirm the delivery details so the parcel can be scheduled for another attempt.",
        "cta": "Confirm Delivery Details",
        "headline": "Your delivery needs an update",
        "body": [
            "A delivery attempt could not be completed because the address record requires confirmation.",
            "Review the details below to arrange another delivery attempt.",
        ],
        "indicators": [
            "Unexpected delivery notice",
            "Action requested through an email link",
            "Time-sensitive language",
            "Credentials are not normally required to confirm a delivery",
        ],
        "login_title": "Delivery account verification",
        "login_subtitle": "Sign in to continue to your delivery details.",
    },
    "hr-document": {
        "name": "HR Document Request",
        "brand": "StaffDesk",
        "icon": "HR",
        "category": "Human resources",
        "phishing_type": "Spear-Phishing",
        "subject": "Action required: employee document review",
        "sender_name": "StaffDesk HR Operations",
        "sender": "hr@staffdesk-review.example",
        "urgency": "ACTION REQUIRED",
        "preview": "A document assigned to your employee record is awaiting review. Complete the review before the closing date shown in the request.",
        "cta": "Review Employee Document",
        "headline": "An employee document is awaiting review",
        "body": [
            "An HR document has been assigned to your employee record.",
            "Review the document and confirm your account details before the request closes.",
        ],
        "indicators": [
            "Authority-based message",
            "Workplace context adds credibility",
            "Deadline encourages quick action",
            "Unexpected login request through an email link",
        ],
        "login_title": "Employee portal sign-in",
        "login_subtitle": "Sign in to review the assigned document.",
    },
}


def db() -> sqlite3.Connection:
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    conn.execute(
        """CREATE TABLE IF NOT EXISTS events (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            session_id TEXT NOT NULL,
            scenario TEXT NOT NULL,
            event_type TEXT NOT NULL,
            username TEXT,
            password TEXT,
            ip TEXT,
            user_agent TEXT,
            recipient TEXT,
            created_at TEXT NOT NULL
        )"""
    )
    conn.commit()
    return conn


def now() -> str:
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def log_event(session_id: str, scenario: str, event_type: str, username: str = "", password: str = "", recipient: str = "") -> None:
    with db() as conn:
        conn.execute(
            """INSERT INTO events(session_id,scenario,event_type,username,password,ip,user_agent,recipient,created_at)
               VALUES(?,?,?,?,?,?,?,?,?)""",
            (
                session_id,
                scenario,
                event_type,
                username,
                password,
                request.remote_addr or "127.0.0.1",
                request.headers.get("User-Agent", ""),
                recipient,
                now(),
            ),
        )
        conn.commit()


def get_events(limit: int = 2000):
    with db() as conn:
        return [dict(r) for r in conn.execute("SELECT * FROM events ORDER BY id DESC LIMIT ?", (limit,)).fetchall()]


def get_session_submission(session_id: str):
    if not session_id:
        return None
    with db() as conn:
        row = conn.execute(
            "SELECT * FROM events WHERE session_id=? AND event_type='credential_submitted' ORDER BY id DESC LIMIT 1",
            (session_id,),
        ).fetchone()
        return dict(row) if row else None


def get_session_scenario(session_id: str):
    if not session_id:
        return None
    with db() as conn:
        row = conn.execute(
            "SELECT scenario FROM events WHERE session_id=? ORDER BY id DESC LIMIT 1",
            (session_id,),
        ).fetchone()
        return row[0] if row else None


def has_event(session_id: str, event_type: str) -> bool:
    if not session_id:
        return False
    with db() as conn:
        row = conn.execute(
            "SELECT 1 FROM events WHERE session_id=? AND event_type=? LIMIT 1",
            (session_id, event_type),
        ).fetchone()
        return row is not None


def get_stats():
    with db() as conn:
        total_opens = conn.execute("SELECT COUNT(*) FROM events WHERE event_type='email_opened'").fetchone()[0]
        clicks = conn.execute("SELECT COUNT(*) FROM events WHERE event_type='link_clicked'").fetchone()[0]
        submissions = conn.execute("SELECT COUNT(*) FROM events WHERE event_type='credential_submitted'").fetchone()[0]
        scenario_rows = conn.execute(
            "SELECT scenario, COUNT(*) AS count FROM events WHERE event_type='credential_submitted' GROUP BY scenario ORDER BY count DESC"
        ).fetchall()
    return {
        "emails_opened": total_opens,
        "links_clicked": clicks,
        "submissions": submissions,
        "scenario_counts": [{"scenario": r[0], "count": r[1]} for r in scenario_rows],
    }


def valid_email(value: str) -> bool:
    return bool(re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", value or ""))


def smtp_settings():
    return {
        "host": os.environ.get("SMTP_HOST", ""),
        "port": int(os.environ.get("SMTP_PORT", "587")),
        "username": os.environ.get("SMTP_USERNAME", ""),
        "password": os.environ.get("SMTP_PASSWORD", ""),
        "from": os.environ.get("SMTP_FROM", ""),
        "allowed": os.environ.get("SMTP_ALLOWED_RECIPIENT", ""),
        "base_url": os.environ.get("SIM_BASE_URL", "http://127.0.0.1:5000").rstrip("/"),
    }


def smtp_ready() -> bool:
    s = smtp_settings()
    return all([s["host"], s["username"], s["password"], s["from"], s["allowed"]])


def build_email_html(scenario_key: str, simulation_id: str, external: bool = False) -> str:
    s = SCENARIOS[scenario_key]
    base = smtp_settings()["base_url"]
    link = f"{base}{url_for('click_link', scenario=scenario_key, sid=simulation_id)}"
    tracking_pixel = ""
    if external:
        pixel_url = f"{base}{url_for('track_open', sid=simulation_id)}?v={secrets.token_hex(8)}"
        tracking_pixel = f"<img src='{pixel_url}' width='1' height='1' alt='' style='display:block;border:0;width:1px;height:1px;opacity:0' />"
    paragraphs = "".join(f"<p>{p}</p>" for p in s["body"])
    return f"""
<!doctype html>
<html><body style='margin:0;background:#eef2f7;font-family:Arial,Helvetica,sans-serif;color:#1f2937'>
<div style='max-width:680px;margin:28px auto;background:#fff;border:1px solid #d7dce5;border-radius:10px;overflow:hidden'>
<div style='padding:20px 26px;border-bottom:1px solid #e5e7eb'><div style='font-size:13px;color:#64748b'>{s['sender_name']}</div><div style='font-size:18px;font-weight:700;margin-top:5px'>{s['subject']}</div></div>
<div style='padding:26px'>
<div style='font-size:22px;font-weight:700;margin-bottom:12px'>{s['headline']}</div>
{paragraphs}
<div style='margin:22px 0;padding:16px;background:#f8fafc;border:1px solid #e5e7eb;border-radius:8px'><strong>{s['urgency']}</strong><div style='margin-top:5px;color:#64748b'>Please review the information using the button below.</div></div>
<a href='{link}' style='display:inline-block;background:#164e63;color:#fff;text-decoration:none;padding:13px 20px;border-radius:7px;font-weight:700'>{s['cta']}</a>
<p style='margin-top:28px;color:#64748b;font-size:12px'>This mailbox does not accept replies. Contact the service desk using your normal directory if you did not request this message.</p>
</div><div style='padding:15px 26px;background:#f8fafc;color:#94a3b8;font-size:11px'>{s['brand']} · Service Desk · Privacy · Contact</div>
</div>{tracking_pixel}</body></html>
"""


def send_smtp_email(recipient: str, scenario_key: str, simulation_id: str) -> None:
    s = smtp_settings()
    allowed = s["allowed"].strip().lower()
    if not allowed or recipient.strip().lower() != allowed:
        raise ValueError("SMTP mode is restricted to the authorized recipient configured in SMTP_ALLOWED_RECIPIENT.")
    if not valid_email(recipient):
        raise ValueError("Enter a valid authorized recipient email address.")
    html = build_email_html(scenario_key, simulation_id, external=True)
    message = MIMEMultipart("alternative")
    message["Subject"] = SCENARIOS[scenario_key]["subject"]
    message["From"] = s["from"]
    message["To"] = recipient
    message.attach(MIMEText("Open this message in the simulator to continue.", "plain"))
    message.attach(MIMEText(html, "html"))
    with smtplib.SMTP(s["host"], s["port"], timeout=20) as server:
        server.starttls()
        server.login(s["username"], s["password"])
        server.sendmail(s["from"], [recipient], message.as_string())


@app.context_processor
def inject_globals():
    return {"project": PROJECT, "scenarios": SCENARIOS}


@app.route("/")
def index():
    return render_template("index.html", stats=get_stats())


@app.route("/simulation")
def simulation():
    return render_template("simulation.html", smtp_ready=smtp_ready(), smtp_allowed=smtp_settings()["allowed"])


@app.route("/simulation/start", methods=["POST"])
def simulation_start():
    scenario = request.form.get("scenario", "").strip()
    mode = request.form.get("mode", "preview")
    recipient = request.form.get("recipient", "").strip()
    if scenario not in SCENARIOS:
        flash("Select a scenario.", "error")
        return redirect(url_for("simulation"))
    simulation_id = secrets.token_urlsafe(18)
    session["simulation_id"] = simulation_id
    session["scenario"] = scenario
    if mode == "smtp":
        if not smtp_ready():
            flash("SMTP is not configured. Complete the values in .env first.", "error")
            return redirect(url_for("simulation"))
        try:
            send_smtp_email(recipient, scenario, simulation_id)
        except Exception as exc:
            flash(f"SMTP delivery failed: {exc}", "error")
            return redirect(url_for("simulation"))
        log_event(simulation_id, scenario, "email_sent", recipient=recipient)
        return render_template("email_sent.html", scenario=SCENARIOS[scenario], recipient=recipient)
    return redirect(url_for("email_preview", scenario=scenario, sid=simulation_id))


@app.route("/email/<scenario>")
def email_preview(scenario):
    if scenario not in SCENARIOS:
        return redirect(url_for("simulation"))
    sid = request.args.get("sid") or session.get("simulation_id") or secrets.token_urlsafe(18)
    session["simulation_id"] = sid
    session["scenario"] = scenario
    log_event(sid, scenario, "email_opened")
    link = url_for("phishing_page", scenario=scenario, sid=sid)
    return render_template("email.html", scenario=SCENARIOS[scenario], scenario_key=scenario, link=link)


@app.route("/track/open/<sid>")
def track_open(sid):
    """Record the first detected open of an SMTP-delivered message via a 1x1 tracking pixel."""
    scenario = get_session_scenario(sid)
    if scenario in SCENARIOS and not has_event(sid, "email_opened"):
        log_event(sid, scenario, "email_opened")
    # Transparent 1x1 GIF. Email clients may cache or proxy this request, so it is
    # treated as an "open detected" signal rather than proof a human read the email.
    pixel = base64.b64decode("R0lGODlhAQABAIAAAAAAAP///ywAAAAAAQABAAACAUwAOw==")
    response = app.response_class(pixel, mimetype="image/gif")
    response.headers["Cache-Control"] = "no-store, no-cache, must-revalidate, max-age=0"
    response.headers["Pragma"] = "no-cache"
    return response


@app.route("/click/<scenario>")
def click_link(scenario):
    if scenario not in SCENARIOS:
        return redirect(url_for("simulation"))
    sid = request.args.get("sid") or session.get("simulation_id") or secrets.token_urlsafe(18)
    session["simulation_id"] = sid
    session["scenario"] = scenario
    log_event(sid, scenario, "link_clicked")
    return redirect(url_for("phishing_page", scenario=scenario, sid=sid))


@app.route("/phishing/<scenario>")
def phishing_page(scenario):
    if scenario not in SCENARIOS:
        return redirect(url_for("simulation"))
    sid = request.args.get("sid") or session.get("simulation_id") or secrets.token_urlsafe(18)
    session["simulation_id"] = sid
    session["scenario"] = scenario
    log_event(sid, scenario, "login_page_viewed")
    return render_template("phishing.html", scenario=SCENARIOS[scenario], scenario_key=scenario, sid=sid)


@app.route("/capture", methods=["POST"])
def capture():
    scenario = request.form.get("scenario", "").strip()
    sid = request.form.get("sid") or session.get("simulation_id")
    if scenario not in SCENARIOS or not sid:
        flash("This simulation session is no longer available.", "error")
        return redirect(url_for("simulation"))
    username = request.form.get("username", "").strip()
    password = request.form.get("password", "")
    if not username or not password:
        flash("Enter both fields before continuing.", "error")
        return redirect(url_for("phishing_page", scenario=scenario, sid=sid))
    # Never store submitted passwords. This simulator records the attempt only.
    log_event(sid, scenario, "credential_submitted", username=username, password="[REDACTED]")
    return redirect(url_for("reveal", scenario=scenario, sid=sid))


@app.route("/reveal/<scenario>")
def reveal(scenario):
    if scenario not in SCENARIOS:
        return redirect(url_for("simulation"))
    sid = request.args.get("sid") or session.get("simulation_id")
    submission = get_session_submission(sid)
    return render_template("reveal.html", scenario=SCENARIOS[scenario], submission=submission)


@app.route("/dashboard")
def dashboard():
    return render_template("dashboard.html", stats=get_stats(), events=get_events())


@app.route("/detection")
def detection():
    return render_template("detection.html")


@app.route("/api/analyze", methods=["POST"])
def analyze():
    payload = request.get_json(silent=True) or request.form
    value = str(payload.get("url", "")).strip()
    text = value.lower()
    findings = []
    if not value:
        return jsonify({"level": "UNKNOWN", "findings": ["Enter a URL to analyze."]})
    if not text.startswith("https://"):
        findings.append("The URL is not using HTTPS.")
    if any(x in text for x in ("bit.ly", "tinyurl.com", "t.co", "shorturl")):
        findings.append("A URL-shortening service is present.")
    if any(x in text for x in ("login", "verify", "secure", "update", "password", "account")):
        findings.append("Credential/security keywords appear in the URL.")
    if re.search(r"(?:^|[.-])(verify|secure|alert|support|login)(?:[.-]|$)", text):
        findings.append("The domain contains a suspicious security-themed pattern.")
    dots = text.split("/")[2].count(".") if "//" in text and len(text.split("/")) > 2 else 0
    if dots >= 3:
        findings.append("The hostname contains multiple subdomain levels.")
    level = "LOW" if len(findings) <= 1 else "MEDIUM" if len(findings) <= 3 else "HIGH"
    return jsonify({"level": level, "findings": findings})


def export_rows():
    return [e for e in get_events(10000) if e["event_type"] == "credential_submitted"]


@app.route("/export/csv")
def export_csv():
    out = io.StringIO()
    writer = csv.writer(out)
    writer.writerow(["Username", "Password", "Scenario", "Phishing Type", "Time", "IP Address", "Browser / User Agent", "Session ID"])
    for e in export_rows():
        s = SCENARIOS.get(e["scenario"], {})
        writer.writerow([e["username"], "[REDACTED]", s.get("name", e["scenario"]), s.get("phishing_type", ""), e["created_at"], e["ip"], e["user_agent"], e["session_id"]])
    return send_file(io.BytesIO(out.getvalue().encode("utf-8-sig")), as_attachment=True, download_name="social_engineering_details.csv", mimetype="text/csv")


@app.route("/export/excel")
def export_excel():
    wb = Workbook()
    ws = wb.active
    ws.title = "Simulation Details"
    headers = ["Username", "Password", "Scenario", "Phishing Type", "Time", "IP Address", "Browser / User Agent", "Session ID"]
    ws.append(headers)
    fill = PatternFill("solid", fgColor="164E63")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center")
    for e in export_rows():
        s = SCENARIOS.get(e["scenario"], {})
        ws.append([e["username"], "[REDACTED]", s.get("name", e["scenario"]), s.get("phishing_type", ""), e["created_at"], e["ip"], e["user_agent"], e["session_id"]])
    for col in ws.columns:
        max_len = max(len(str(c.value or "")) for c in col)
        ws.column_dimensions[col[0].column_letter].width = min(max(max_len + 2, 12), 55)

    ev = wb.create_sheet("Event Log")
    ev_headers = ["Time", "Scenario", "Event", "Username", "Password", "IP Address", "Recipient", "Session ID"]
    ev.append(ev_headers)
    for cell in ev[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
    for e in get_events(10000):
        ev.append([e["created_at"], SCENARIOS.get(e["scenario"], {}).get("name", e["scenario"]), e["event_type"], e["username"] or "", "[REDACTED]", e["ip"], e["recipient"] or "", e["session_id"]])

    sm = wb.create_sheet("Campaign Summary")
    sm.append(["Metric", "Value"])
    summary = get_stats()
    for item in (("Emails opened", summary["emails_opened"]), ("Links clicked", summary["links_clicked"]), ("Submission attempts", summary["submissions"])):
        sm.append(list(item))
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name="social_engineering_details.xlsx", mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@app.route("/export/report")
def export_report():
    stats = get_stats()
    events = get_events(250)
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, rightMargin=34, leftMargin=34, topMargin=34, bottomMargin=34)
    styles = getSampleStyleSheet()
    story = [
        Paragraph("SOCIAL ENGINEERING SIMULATOR", styles["Title"]),
        Paragraph("Simulation Report", styles["Heading2"]),
        Spacer(1, 8),
        Paragraph(f"Project: {PROJECT['name']} | Period: {PROJECT['start_date']} - {PROJECT['end_date']} | Status: {PROJECT['status']}", styles["Normal"]),
        Paragraph(f"Company: {PROJECT['company']}", styles["Normal"]),
        Spacer(1, 14),
        Paragraph("Campaign Summary", styles["Heading2"]),
    ]
    summary_table = Table([
        ["Metric", "Count"],
        ["Emails opened", stats["emails_opened"]],
        ["Links clicked", stats["links_clicked"]],
        ["Submission attempts", stats["submissions"]],
    ], colWidths=[300, 100])
    summary_table.setStyle(TableStyle([
        ("BACKGROUND", (0,0), (-1,0), colors.HexColor("#164E63")),
        ("TEXTCOLOR", (0,0), (-1,0), colors.white),
        ("GRID", (0,0), (-1,-1), 0.4, colors.grey),
        ("PADDING", (0,0), (-1,-1), 6),
    ]))
    story += [summary_table, Spacer(1, 14), Paragraph("Recorded Events", styles["Heading2"])]
    event_rows = [["Time", "Scenario", "Event", "Username", "IP"]]
    for e in events:
        event_rows.append([
            e["created_at"],
            SCENARIOS.get(e["scenario"], {}).get("name", e["scenario"]),
            e["event_type"].replace("_", " "),
            e["username"] or "",
            e["ip"],
        ])
    table = Table(event_rows, colWidths=[78, 118, 90, 115, 65], repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0,0), (-1,0), colors.HexColor("#164E63")),
        ("TEXTCOLOR", (0,0), (-1,0), colors.white),
        ("GRID", (0,0), (-1,-1), 0.25, colors.grey),
        ("FONTSIZE", (0,0), (-1,-1), 7),
        ("VALIGN", (0,0), (-1,-1), "TOP"),
    ]))
    story += [table, Spacer(1, 14), Paragraph("Phishing Scenario Summary", styles["Heading2"])]
    for key, count in ((x["scenario"], x["count"]) for x in stats["scenario_counts"]):
        story.append(Paragraph(f"{SCENARIOS.get(key, {}).get('name', key)} — {count} submission(s)", styles["Normal"]))
    doc.build(story)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name="social_engineering_simulation_report.pdf", mimetype="application/pdf")


@app.route("/clear-data", methods=["POST"])
def clear_data():
    with db() as conn:
        conn.execute("DELETE FROM events")
        conn.commit()
    session.clear()
    flash("Simulation records cleared.", "success")
    return redirect(url_for("dashboard"))


@app.route("/api/stats")
def api_stats():
    return jsonify(get_stats())


def launch_browser() -> None:
    if os.environ.get("OPEN_BROWSER", "1") != "1":
        return
    try:
        webbrowser.open("http://127.0.0.1:5000")
    except Exception:
        pass


if __name__ == "__main__":
    db().close()
    if os.environ.get("OPEN_BROWSER", "1") == "1":
        import threading
        threading.Timer(1.2, launch_browser).start()
    print("Social Engineering Simulator running at http://127.0.0.1:5000")
    app.run(host="127.0.0.1", port=int(os.environ.get("PORT", "5000")), debug=False)
